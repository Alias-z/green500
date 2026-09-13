"""Collect and match current Science Based Targets initiative records."""

from __future__ import annotations

import math
import re
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from pathlib import PurePosixPath
from urllib.parse import urlsplit

import scrapy
from openpyxl import load_workbook
from scrapy.crawler import CrawlerProcess

from green500 import db
from green500.config import load_settings
from green500.storage import validate_public_url

SOURCE_KEY = "sbti"
TARGET_DASHBOARD_URL = "https://sciencebasedtargets.org/target-dashboard"
COMPANIES_FILENAME = "companies-excel.xlsx"
TARGETS_FILENAME = "targets-excel.xlsx"
VERIFIED_WORKBOOK_URLS = {
    "companies": (
        "https://files.sciencebasedtargets.org/production/files/companies-excel.xlsx"
    ),
    "targets": (
        "https://files.sciencebasedtargets.org/production/files/targets-excel.xlsx"
    ),
}

COMPANY_REQUIRED_COLUMNS = {
    "sbti_id",
    "company_name",
    "near_term_status",
    "net_zero_status",
    "date_updated",
}
TARGET_REQUIRED_COLUMNS = {
    "row_entry_id",
    "sbti_id",
    "company_name",
    "status",
    "target_wording",
    "date_published",
}


def _json_value(value):
    """Convert typed spreadsheet values to stable JSON values."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        if not math.isfinite(value):
            raise ValueError("SBTi workbook contains a non-finite number.")
        if value.is_integer():
            return int(value)
    return value


def normalize_sbti_id(value) -> str:
    """Normalize an Excel numeric identifier without changing its digits."""
    if value is None or isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            return ""
        return str(int(value))
    text = str(value).strip()
    if not text:
        return ""
    try:
        number = Decimal(text)
    except InvalidOperation:
        return text
    if not number.is_finite() or number != number.to_integral_value():
        return ""
    return format(number.quantize(Decimal(1)), "f")


def parse_workbook_rows(
    body: bytes,
    *,
    sheet_name: str,
    required_columns: set[str],
) -> list[dict]:
    """Parse one typed worksheet and retain its one-based Excel row numbers."""
    if not body.startswith(b"PK"):
        raise ValueError("SBTi workbook response is not an XLSX file.")
    try:
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("SBTi workbook could not be opened as XLSX.") from error
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"SBTi workbook is missing the {sheet_name!r} sheet.")
        worksheet = workbook[sheet_name]
        source_rows = worksheet.iter_rows(values_only=True)
        try:
            raw_headers = next(source_rows)
        except StopIteration as error:
            raise ValueError("SBTi workbook has no header row.") from error
        headers = [str(value).strip() if value is not None else None for value in raw_headers]
        named_headers = [value for value in headers if value]
        if len(named_headers) != len(set(named_headers)):
            raise ValueError("SBTi workbook contains duplicate named columns.")
        missing = sorted(required_columns - set(named_headers))
        if missing:
            raise ValueError(
                "SBTi workbook is missing required columns: " + ", ".join(missing)
            )
        rows = []
        for source_row, values in enumerate(source_rows, start=2):
            fields = {
                header: _json_value(values[index] if index < len(values) else None)
                for index, header in enumerate(headers)
                if header
            }
            if not any(value not in (None, "") for value in fields.values()):
                continue
            sbti_id = normalize_sbti_id(fields.get("sbti_id"))
            if not sbti_id or not str(fields.get("company_name") or "").strip():
                raise ValueError(
                    f"SBTi worksheet row {source_row} is missing its company identity."
                )
            fields["sbti_id"] = sbti_id
            fields["source_row"] = source_row
            rows.append(fields)
        if not rows:
            raise ValueError("SBTi workbook contains no data rows.")
        return rows
    finally:
        workbook.close()


_SHARE_CLASS_SUFFIX = re.compile(
    r"(?:\s+|\s*\(\s*)(?:class|cl)\s+[a-z0-9]+\s*\)?\s*$",
    re.IGNORECASE,
)


def normalize_company_name(name: str) -> str:
    """Normalize only case, punctuation, spacing and a trailing share class."""
    text = unicodedata.normalize("NFKC", str(name or "")).strip()
    text = _SHARE_CLASS_SUFFIX.sub("", text)
    text = "".join(
        character
        for character in text.casefold()
        if not unicodedata.category(character).startswith("P")
    )
    return " ".join(text.split())


def _normalized_overrides(raw_overrides: dict | None) -> dict[str, str]:
    """Validate optional explicit CIK-to-SBTi identifier mappings."""
    result = {}
    for raw_cik, raw_sbti_id in (raw_overrides or {}).items():
        cik = str(raw_cik).strip()
        if cik.isdigit() and len(cik) <= 10:
            cik = cik.zfill(10)
        if not re.fullmatch(r"\d{10}", cik):
            raise ValueError("Every SBTi mapping override key must be a valid CIK.")
        sbti_id = normalize_sbti_id(raw_sbti_id)
        if not sbti_id:
            raise ValueError("Every SBTi mapping override must contain an SBTi identifier.")
        result[cik] = sbti_id
    return result


def match_current_companies(
    current_companies: list[dict],
    sbti_companies: list[dict],
    *,
    sbti_id_overrides: dict | None = None,
) -> dict[str, list[dict]]:
    """Match only unique normalized names, with explicit identifier overrides."""
    source_by_name = defaultdict(list)
    source_by_id = defaultdict(list)
    for row in sbti_companies:
        source_by_name[normalize_company_name(row["company_name"])].append(row)
        source_by_id[normalize_sbti_id(row["sbti_id"])].append(row)

    current_name_counts = Counter(
        normalize_company_name(company["name"]) for company in current_companies
    )
    overrides = _normalized_overrides(sbti_id_overrides)
    result: dict[str, list[dict]] = {
        "matched": [],
        "unmatched": [],
        "ambiguous": [],
    }
    for company in current_companies:
        cik = str(company["cik"])
        normalized_name = normalize_company_name(company["name"])
        if cik in overrides:
            candidates = source_by_id.get(overrides[cik], [])
            method = "sbti_id_override"
            reason = "The explicit SBTi identifier mapping did not resolve to one source row."
        else:
            candidates = source_by_name.get(normalized_name, [])
            method = "unique_name"
            reason = "The normalized company name did not resolve to one source row."

        summary = {
            "cik": cik,
            "company_name": company["name"],
            "normalized_name": normalized_name,
            "match_method": method,
        }
        if not candidates:
            result["unmatched"].append({**summary, "reason": reason})
            continue
        if len(candidates) != 1 or (
            method == "unique_name" and current_name_counts[normalized_name] != 1
        ):
            result["ambiguous"].append(
                {
                    **summary,
                    "reason": reason,
                    "current_company_name_match_count": current_name_counts[
                        normalized_name
                    ],
                    "source_name_match_count": len(candidates),
                    "candidate_sbti_ids": sorted(
                        {normalize_sbti_id(row["sbti_id"]) for row in candidates}
                    ),
                }
            )
            continue
        result["matched"].append(
            {
                **summary,
                "source_company": candidates[0],
            }
        )
    return result


def status_observations(source_company: dict) -> list[dict]:
    """Create observations only for status values actually present in the workbook."""
    published_at = source_company.get("date_updated") or None
    source_fields = {
        key: value for key, value in source_company.items() if key != "source_row"
    }
    source_record = {
        "workbook": COMPANIES_FILENAME,
        "source_row": source_company["source_row"],
        "sbti_id": source_company["sbti_id"],
        "company_name": source_company["company_name"],
        "original_status": {
            "near_term_status": source_company.get("near_term_status"),
            "long_term_status": source_company.get("long_term_status"),
            "net_zero_status": source_company.get("net_zero_status"),
        },
        "fields": source_fields,
    }
    observations = []
    for metric_code, column in (
        ("near_term_target_status", "near_term_status"),
        ("long_term_target_status", "long_term_status"),
        ("net_zero_target_status", "net_zero_status"),
    ):
        value = source_company.get(column)
        if value in (None, ""):
            continue
        observations.append(
            {
                "metric_code": metric_code,
                "value": value,
                "unit": None,
                "period_start": None,
                "period_end": None,
                "period_type": "current",
                "published_at": published_at,
                "source_record": source_record,
            }
        )
    return observations


def climate_targets_observation(
    source_company: dict, target_rows: list[dict]
) -> dict | None:
    """Preserve every target row together; an absent row does not mean no target."""
    if not target_rows:
        return None
    ordered_rows = sorted(
        target_rows,
        key=lambda row: (str(row.get("row_entry_id") or ""), row["source_row"]),
    )
    source_rows = [
        {
            "source_row": row["source_row"],
            "row_entry_id": row.get("row_entry_id"),
            "sbti_id": row["sbti_id"],
            "status": row.get("status"),
            "date_published": row.get("date_published"),
        }
        for row in ordered_rows
    ]
    return {
        "metric_code": "climate_targets",
        "value": ordered_rows,
        "unit": None,
        "period_start": None,
        "period_end": None,
        "period_type": "current",
        "published_at": source_company.get("date_updated") or None,
        "source_record": {
            "workbook": TARGETS_FILENAME,
            "company_workbook_source_row": source_company["source_row"],
            "sbti_id": source_company["sbti_id"],
            "company_name": source_company["company_name"],
            "target_rows": source_rows,
        },
    }


def resolve_workbook_links(response) -> dict[str, str]:
    """Resolve the two current workbook links observed on the SBTi dashboard."""
    links = {}
    expected = {
        COMPANIES_FILENAME: "companies",
        TARGETS_FILENAME: "targets",
    }
    for href in response.css("a::attr(href)").getall():
        url = response.urljoin(href)
        filename = PurePosixPath(urlsplit(url).path).name.casefold()
        key = expected.get(filename)
        if key and key not in links:
            links[key] = validate_public_url(url)
    return links


class SbtiSpider(scrapy.Spider):
    """Fetch the SBTi dashboard and its two published current-data workbooks."""

    name = "green500_sbti"

    def __init__(self, dashboard_url=TARGET_DASHBOARD_URL, **kwargs):
        super().__init__(**kwargs)
        self.dashboard_url = validate_public_url(dashboard_url)
        self.downloads: dict[str, dict] = {}
        self.failures: list[dict] = []
        self.resolved_urls: dict[str, str] = {}
        self.link_sources: dict[str, str] = {}

    async def start(self):
        """Begin with the stable dashboard so annual link changes remain observable."""
        yield scrapy.Request(
            self.dashboard_url,
            callback=self.parse_dashboard,
            errback=self.record_failure,
        )

    def parse_dashboard(self, response):
        """Prefer observed workbook links and use verified URLs only when absent."""
        observed = resolve_workbook_links(response)
        for key, verified_url in VERIFIED_WORKBOOK_URLS.items():
            url = observed.get(key, verified_url)
            self.resolved_urls[key] = url
            self.link_sources[key] = "dashboard" if key in observed else "verified_fallback"
            yield scrapy.Request(
                url,
                callback=self.parse_workbook,
                errback=self.record_failure,
                cb_kwargs={"workbook_key": key},
            )

    def parse_workbook(self, response, workbook_key: str):
        """Retain exact response bytes and response metadata for later persistence."""
        self.downloads[workbook_key] = {
            "url": response.url,
            "body": response.body,
            "content_type": response.headers.get("Content-Type", b"").decode("latin1")
            or "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        }

    def record_failure(self, failure):
        """Retain a bounded error for every failed dashboard or workbook request."""
        self.failures.append(
            {
                "url": failure.request.url,
                "error": failure.getErrorMessage()[:500],
            }
        )


def _download_sbti_workbooks(settings, dashboard_url: str) -> dict[str, dict]:
    """Run one bounded Scrapy process and return both exact workbook responses."""
    process = CrawlerProcess(
        {
            "USER_AGENT": "Green500/0.1 (public environmental research)",
            "ROBOTSTXT_OBEY": True,
            "CONCURRENT_REQUESTS": 2,
            "CONCURRENT_REQUESTS_PER_DOMAIN": 2,
            "DOWNLOAD_DELAY": 0.25,
            "DOWNLOAD_TIMEOUT": 45,
            "DOWNLOAD_MAXSIZE": settings.max_document_bytes,
            "DOWNLOAD_WARNSIZE": settings.max_document_bytes,
            "DOWNLOAD_FAIL_ON_DATALOSS": True,
            "RETRY_TIMES": 2,
            "REDIRECT_MAX_TIMES": 5,
            "LOG_LEVEL": "WARNING",
            "TELNETCONSOLE_ENABLED": False,
            "TWISTED_REACTOR": "twisted.internet.asyncioreactor.AsyncioSelectorReactor",
        }
    )
    crawler = process.create_crawler(SbtiSpider)
    process.crawl(crawler, dashboard_url=dashboard_url)
    process.start()
    spider = crawler.spider
    if not spider:
        raise RuntimeError("SBTi collector did not start.")
    if spider.failures:
        errors = "; ".join(item["error"] for item in spider.failures)[:1000]
        raise RuntimeError("SBTi download failed: " + errors)
    if set(spider.downloads) != {"companies", "targets"}:
        raise RuntimeError("SBTi collector did not download both current workbooks.")
    return spider.downloads


def _current_companies(settings, company_cik: str | None = None) -> list[dict]:
    """Read every current company, or the one company named by a refresh task."""
    with db.connect(settings) as conn:
        return conn.execute(
            """SELECT cik,name FROM companies
            WHERE is_current AND (%s::text IS NULL OR cik=%s)
            ORDER BY cik""",
            (company_cik, company_cik),
        ).fetchall()


def _record_global_failure(data_store, settings, companies, task, error, dataset_id=None):
    """Record a source-wide failure for every current company before raising it."""
    for company in companies:
        data_store.record_check(
            settings,
            source_key=SOURCE_KEY,
            company_cik=company["cik"],
            status="failed",
            error=str(error)[:1000],
            dataset_id=dataset_id,
            details={"failure_scope": "source_collection_or_parsing"},
        )


def _dataset_result(receipt: dict) -> dict:
    """Return a task-result receipt whose timestamp can be stored as JSON."""
    return {
        key: value.isoformat() if isinstance(value, (date, datetime)) else value
        for key, value in receipt.items()
    }


def collect_targets_task(task: dict) -> dict:
    """Download, version, conservatively match and save current SBTi targets."""
    from green500 import data_store

    settings = load_settings()
    current_companies = _current_companies(settings, task.get("company_cik"))
    if not current_companies:
        raise ValueError("No current companies are available for SBTi matching.")
    task_input = task.get("input") or {}
    dashboard_url = validate_public_url(
        task_input.get("dashboard_url", TARGET_DASHBOARD_URL)
    )
    try:
        downloads = _download_sbti_workbooks(settings, dashboard_url)
    except Exception as error:
        _record_global_failure(data_store, settings, current_companies, task, error)
        raise

    companies_dataset = None
    try:
        companies_dataset = data_store.save_dataset(
            settings,
            source_key=SOURCE_KEY,
            url=downloads["companies"]["url"],
            body=downloads["companies"]["body"],
            content_type=downloads["companies"]["content_type"],
            task_id=task.get("id"),
        )
        targets_dataset = data_store.save_dataset(
            settings,
            source_key=SOURCE_KEY,
            url=downloads["targets"]["url"],
            body=downloads["targets"]["body"],
            content_type=downloads["targets"]["content_type"],
            task_id=task.get("id"),
        )
        company_rows = parse_workbook_rows(
            downloads["companies"]["body"],
            sheet_name="Data",
            required_columns=COMPANY_REQUIRED_COLUMNS,
        )
        target_rows = parse_workbook_rows(
            downloads["targets"]["body"],
            sheet_name="WebsiteData",
            required_columns=TARGET_REQUIRED_COLUMNS,
        )
    except Exception as error:
        _record_global_failure(
            data_store,
            settings,
            current_companies,
            task,
            error,
            companies_dataset["id"] if companies_dataset else None,
        )
        raise

    matches = match_current_companies(
        current_companies,
        company_rows,
        sbti_id_overrides=task_input.get("sbti_ids"),
    )
    targets_by_sbti_id = defaultdict(list)
    for row in target_rows:
        targets_by_sbti_id[row["sbti_id"]].append(row)

    for outcome_name in ("unmatched", "ambiguous"):
        for outcome in matches[outcome_name]:
            data_store.record_check(
                settings,
                source_key=SOURCE_KEY,
                company_cik=outcome["cik"],
                status=outcome_name,
                dataset_id=companies_dataset["id"],
                details={
                    key: value
                    for key, value in outcome.items()
                    if key not in {"cik"}
                },
            )

    failed = []
    for match in matches["matched"]:
        company_cik = match["cik"]
        source_company = match["source_company"]
        match_method = match["match_method"]
        company_targets = targets_by_sbti_id.get(source_company["sbti_id"], [])
        try:
            statuses = status_observations(source_company)
            saved_observation_count = 0
            if statuses:
                saved_observation_count += data_store.save_observations(
                    settings,
                    source_key=SOURCE_KEY,
                    company_cik=company_cik,
                    dataset_id=companies_dataset["id"],
                    observations=statuses,
                    match_method=match_method,
                    task_id=task.get("id"),
                )
            target_observation = climate_targets_observation(
                source_company, company_targets
            )
            if target_observation:
                saved_observation_count += data_store.save_observations(
                    settings,
                    source_key=SOURCE_KEY,
                    company_cik=company_cik,
                    dataset_id=targets_dataset["id"],
                    observations=[target_observation],
                    match_method=match_method,
                    task_id=task.get("id"),
                )
            data_store.record_check(
                settings,
                source_key=SOURCE_KEY,
                company_cik=company_cik,
                status="succeeded",
                dataset_id=companies_dataset["id"],
                details={
                    "match_method": match_method,
                    "company_name": match["company_name"],
                    "source_company_name": source_company["company_name"],
                    "sbti_id": source_company["sbti_id"],
                    "company_workbook_source_row": source_company["source_row"],
                    "target_row_count": len(company_targets),
                    "saved_observation_count": saved_observation_count,
                    "targets_dataset_id": targets_dataset["id"],
                },
            )
        except Exception as error:  # noqa: BLE001 - retain every company outcome
            failed.append(
                {
                    "cik": company_cik,
                    "company_name": match["company_name"],
                    "error": str(error)[:1000],
                }
            )
            data_store.record_check(
                settings,
                source_key=SOURCE_KEY,
                company_cik=company_cik,
                status="failed",
                error=str(error)[:1000],
                dataset_id=companies_dataset["id"],
                details={
                    "match_method": match_method,
                    "sbti_id": source_company["sbti_id"],
                },
            )

    if not matches["matched"]:
        raise ValueError("SBTi collection matched zero current companies.")
    if len(failed) == len(matches["matched"]):
        raise RuntimeError("Every matched SBTi company failed during persistence.")

    return {
        "source_key": SOURCE_KEY,
        "dashboard_url": dashboard_url,
        "companies_dataset": _dataset_result(companies_dataset),
        "targets_dataset": _dataset_result(targets_dataset),
        "matched": [
            {
                "cik": match["cik"],
                "company_name": match["company_name"],
                "sbti_id": match["source_company"]["sbti_id"],
                "match_method": match["match_method"],
            }
            for match in matches["matched"]
        ],
        "unmatched": matches["unmatched"],
        "ambiguous": matches["ambiguous"],
        "failed": failed,
    }
